"""Crea el informe ICONIX desde su especificación; conserva dos carátulas.

Ejecutar con el Python empaquetado para documentos. Actualizar campos y renderizar
con render_iconix_document.ps1 después de crear el Word.
"""
import json
import re
import unicodedata
from copy import deepcopy
from hashlib import sha256
from pathlib import Path
from collections import defaultdict
from shutil import copy2
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from iconix_diagrams import Drawing, graph_boxes, robustness, sequence, use_cases, state_diagram, qr_code

ROOT=Path(__file__).resolve().parents[1]
SOURCE=ROOT/'docs/Plataforma_Vestidor3D_Grupo18.docx'
BASE=ROOT/'docs/saas/iconix'
OUTPUT=ROOT/'docs/Lumina_SaaS_ICONIX_Grupo18.docx'
ART=ROOT/'.local/iconix/figures'
WIDTH=7.08
SPEC=json.loads((BASE/'especificacion.json').read_text(encoding='utf8'))
TABLES={t['name']:t for t in SPEC['data']['existing_tables']+SPEC['data']['new_tables']}
for change in SPEC['data']['changes']:
    for field in change['fields']:
        if not any(f['name']==field['name'] for f in TABLES[change['table']]['fields']):
            TABLES[change['table']]['fields'].append({**field,'planned':True})

ALIASES={
 'Tienda':'platform.tenant','Plan':'platform.plan','Suscripcion':'platform.subscription','PagoSuscripcion':'platform.billing_event',
 'Usuario':'usuario','Sesion':'sesion','Rol':'rol','AsignacionUbicacion':'usuario_ubicacion','Ubicacion':'ubicacion',
 'Producto':'producto','Variante':'variante','Direccion':'direccion',
 'Inventario':'inventario','MovimientoStock':'movimiento_stock','OperacionSincronizacion':'operacion_sincronizacion',
 'CapturaCorporal':'captura_corporal','TrabajoAvatar':'trabajo_avatar','Avatar':'avatar','SesionVestidor':'sesion_vestidor',
 'Carrito':'carrito','Pedido':'pedido','Pago':'pago','HistorialPedido':'historial_pedido','Devolucion':'devolucion','Reembolso':'reembolso',
 'ComprobanteOperacion':'operacion_sincronizacion','DefinicionReporte':'definicion_reporte','ConjuntoReporte':'conjunto_reporte',
 'ConsultaAnalitica':'consulta_analitica','ConjuntoEntrenamiento':'conjunto_entrenamiento','ModeloAnalitico':'modelo_analitico',
 'EjecucionAnalitica':'ejecucion_analitica','Recomendacion':'recomendacion',
}
LOCAL_ATTRIBUTES={
 'RecursoVestidor':['id : UUID','varianteId : UUID','formato : GLB o PNG','sha256 : texto'],
 'PaqueteOffline':['id : UUID','propietaria : UUID','tienda : slug','estado : parcial/completo'],
 'OperacionOffline':['operationId : UUID','deviceId : UUID','baseVersion : entero','estado : EstadoCola'],
 'RecursoLocal':['claveLocal : texto','sha256 : texto','bytes : entero','verificado : booleano'],
 'ConflictoSincronizacion':['operacionId : UUID','versionActual : entero','intencion : JSON','resolucion : JSON'],
 'DisenoDashboard':['version : entero','paneles : JSON','filtros : JSON','seleccion : JSON'],
 'ResultadoAnalitico':['ejecucionId : UUID','procedencia : texto','metricas : JSON','filas : JSON'],
}
UI_ATTRS={
 'CU01':['nombre','slug','correoPropietaria','plan'], 'CU02':['planElegido','importeBOB','proveedor','confirmacion'],
 'CU03':['tienda','correo','credencialTemporal'], 'CU04':['persona','roles','ubicaciones'], 'CU05':['nombre','tipo','padreId'],
 'CU06':['prenda','variantes','recursos','versionBase'], 'CU07':['busqueda','talla','categoria','ubicacion'],
 'CU08':['perfil','direccion','telefono'], 'CU09':['ubicacion','variante','delta','motivo'],
 'CU10':['variante','conteoObservado','versionBase'], 'CU11':['origen','destino','cantidad','idempotencia'],
 'CU12':['fotosTemporales','altura','consentimiento'], 'CU13':['avatarSeleccionado','medidas','decision'],
 'CU14':['avatarAprobado','modeloPrenda','vista'], 'CU15':['poseLocal','imagenPrenda','calibracion'],
 'CU16':['carrito','direccion','metodoPago','idempotencia'], 'CU17':['ubicacion','lineasVenta','metodoPago'],
 'CU18':['pedido','estadoSolicitado','version'], 'CU19':['pedido','detalleDevuelto','motivo'],
 'CU20':['gustos','talla','paquete','avance'], 'CU21':['borrador','fotosLocales','operaciones'],
 'CU22':['inventarioLocal','deltaOConteo','pendientes'], 'CU23':['cola','conflictos','conectividad'],
 'CU24':['filtros','metricas','paneles','cobertura'], 'CU25':['textoOAudioTemporal','interpretacion','confirmacion'],
 'CU26':['conjuntoSintetico','parametros','particiones'], 'CU27':['modelo','horizonte','ambito','procedencia'],
}


class Book:
    def __init__(self):
        self.doc=Document(SOURCE);self.counter=0;self.bookmark_id=12000;self.num_id=5000;self.figure_paths=[]
        self.original_sha=sha256(SOURCE.read_bytes()).hexdigest()
        body=self.doc._element.body
        end=deepcopy(self.doc.sections[2]._sectPr)
        for el in list(body)[16:]:body.remove(el)
        body.append(end)
        for index in [3,11]:
            p=self.doc.paragraphs[index]
            p.runs[0].text='Plataforma SaaS de Moda Femenina con\nVestidor 3D y Analítica Comercial'
            for run in p.runs[1:]:run.text=''
            p.style=self.doc.styles['Title']
        for name,size in [('Normal',10.5),('Title',21),('Heading 1',16),('Heading 2',13),('Heading 3',11.5),('Caption',9)]:
            s=self.doc.styles[name];s.font.name='Arial';s.font.size=Pt(size);s.font.color.rgb=RGBColor(0,0,0)
            if s._element.pPr is not None:
                border=s._element.pPr.find(qn('w:pBdr'))
                if border is not None:s._element.pPr.remove(border)
        normal=self.doc.styles['Normal'].paragraph_format
        normal.line_spacing=1.0583;normal.space_after=Pt(5)
        for name in ['Heading 1','Heading 2','Heading 3']:
            s=self.doc.styles[name].paragraph_format;s.keep_with_next=True;s.space_before=Pt(10);s.space_after=Pt(6)
        for name,size in [('TOC 1',10.5),('TOC 2',10)]:
            if name not in self.doc.styles:self.doc.styles.add_style(name,WD_STYLE_TYPE.PARAGRAPH)
            s=self.doc.styles[name];s.font.name='Arial';s.font.size=Pt(size);s.font.color.rgb=RGBColor(0,0,0)
            s.paragraph_format.space_after=Pt(2)
        section=self.doc.sections[-1]
        section.different_first_page_header_footer=False
        section.footer.is_linked_to_previous=False
        footer=section.footer
        for element in list(footer._element):footer._element.remove(element)
        p=footer.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.RIGHT
        field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');p._p.append(field)
        settings=self.doc.settings.element
        update=OxmlElement('w:updateFields');update.set(qn('w:val'),'true');settings.append(update)
        self.make_numbering()

    def make_numbering(self):
        root=self.doc.part.numbering_part.element
        a=OxmlElement('w:abstractNum');a.set(qn('w:abstractNumId'),'5000')
        level=OxmlElement('w:lvl');level.set(qn('w:ilvl'),'0')
        for name,value in [('start','1'),('numFmt','decimal'),('lvlText','%1.'),('suff','space')]:
            e=OxmlElement('w:'+name);e.set(qn('w:val'),value);level.append(e)
        a.append(level);root.append(a)

    def heading(self,text,level=1,newpage=False,anchor=None):
        p=self.doc.add_paragraph(text,style=f'Heading {level}')
        if newpage:p.paragraph_format.page_break_before=True
        if anchor:self.bookmark(p,anchor)
        return p

    def para(self,text='',bold_label=None,small=False,keep=False):
        p=self.doc.add_paragraph();p.paragraph_format.keep_with_next=keep
        if bold_label:p.add_run(bold_label+' ').bold=True
        run=p.add_run(str(text));
        if small:run.font.size=Pt(9)
        return p

    def numbered(self,items):
        self.num_id+=1;root=self.doc.part.numbering_part.element
        num=OxmlElement('w:num');num.set(qn('w:numId'),str(self.num_id))
        ab=OxmlElement('w:abstractNumId');ab.set(qn('w:val'),'5000');num.append(ab);root.append(num)
        override=OxmlElement('w:lvlOverride');override.set(qn('w:ilvl'),'0')
        start=OxmlElement('w:startOverride');start.set(qn('w:val'),'1');override.append(start);num.append(override)
        for item in items:
            p=self.para(item);p.paragraph_format.left_indent=Inches(.20);p.paragraph_format.first_line_indent=Inches(-.20)
            ppr=p._p.get_or_add_pPr();n=OxmlElement('w:numPr')
            for name,value in [('ilvl','0'),('numId',str(self.num_id))]:
                e=OxmlElement('w:'+name);e.set(qn('w:val'),value);n.append(e)
            ppr.append(n)

    def bookmark(self,p,name):
        self.bookmark_id+=1
        start=OxmlElement('w:bookmarkStart');start.set(qn('w:id'),str(self.bookmark_id));start.set(qn('w:name'),name)
        end=OxmlElement('w:bookmarkEnd');end.set(qn('w:id'),str(self.bookmark_id))
        p._p.insert(0,start);p._p.append(end)

    def link(self,p,text,anchor=None,url=None):
        from docx.opc.constants import RELATIONSHIP_TYPE
        link=OxmlElement('w:hyperlink')
        if url:link.set(qn('r:id'),self.doc.part.relate_to(url,RELATIONSHIP_TYPE.HYPERLINK,is_external=True))
        else:link.set(qn('w:anchor'),anchor)
        run=OxmlElement('w:r');props=OxmlElement('w:rPr')
        color=OxmlElement('w:color');color.set(qn('w:val'),'000000');props.append(color)
        under=OxmlElement('w:u');under.set(qn('w:val'),'single');props.append(under);run.append(props)
        t=OxmlElement('w:t');t.set(qn('xml:space'),'preserve');t.text=text;run.append(t);link.append(run);p._p.append(link)

    def table(self,headers,rows,widths=None,font=9.5):
        t=self.doc.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
        widths=widths or [WIDTH/len(headers)]*len(headers)
        for col,width in zip(t.columns,widths):col.width=Inches(width)
        for cell,width,text in zip(t.rows[0].cells,widths,headers):cell.width=Inches(width);cell.text=str(text)
        for row in rows:
            for cell,width,text in zip(t.add_row().cells,widths,row):cell.width=Inches(width);cell.text=str(text)
        for i,row in enumerate(t.rows):
            for j,cell in enumerate(row.cells):
                cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                pr=cell._tc.get_or_add_tcPr()
                margins=OxmlElement('w:tcMar')
                for side in ['top','left','bottom','right']:
                    el=OxmlElement('w:'+side);el.set(qn('w:w'),'75');el.set(qn('w:type'),'dxa');margins.append(el)
                pr.append(margins)
                borders=OxmlElement('w:tcBorders')
                for side in ['top','left','bottom','right']:
                    e=OxmlElement('w:'+side);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'5');e.set(qn('w:color'),'D9D9D9');borders.append(e)
                pr.append(borders)
                if i==0:
                    shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E7E6E6');pr.append(shade)
                for p in cell.paragraphs:
                    p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=1.0
                    if i==0:p.paragraph_format.keep_with_next=True
                    for run in p.runs:run.font.name='Arial';run.font.size=Pt(font);run.bold=i==0
            if i==0:
                repeat=OxmlElement('w:tblHeader');row._tr.get_or_add_trPr().append(repeat)
            no_split=OxmlElement('w:cantSplit');row._tr.get_or_add_trPr().append(no_split)
        self.para('').paragraph_format.space_after=Pt(1)
        return t

    def figure(self,path,title,anchor=None,maxheight=7.55,width=WIDTH):
        from PIL import Image
        with Image.open(path) as img:w,h=img.size
        width=min(width,maxheight*w/h)
        p=self.doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after=Pt(0);p.paragraph_format.keep_with_next=True
        picture=p.add_run().add_picture(str(path),width=Inches(width))
        picture._inline.docPr.set('descr',title)
        self.counter+=1
        self.figure_paths.append(Path(path))
        cap=self.doc.add_paragraph(f'Figura {self.counter}  {title}',style='Caption');cap.alignment=WD_ALIGN_PARAGRAPH.CENTER
        cap.paragraph_format.space_after=Pt(5)
        if anchor:self.bookmark(cap,anchor)
        return self.counter

    def toc(self):
        self.doc.add_paragraph('ÍNDICE',style='Title')
        p=self.doc.add_paragraph();r=p.add_run()
        for kind in ['begin','separate','end']:
            el=OxmlElement('w:fldChar');el.set(qn('w:fldCharType'),kind)
            if kind=='begin':
                r._r.append(el);inst=OxmlElement('w:instrText');inst.set(qn('xml:space'),'preserve');inst.text=' TOC \\o "1-2" \\h \\z \\u ';r._r.append(inst)
            elif kind=='separate':
                r._r.append(el);r.add_text('Índice de secciones y casos de uso')
            else:r._r.append(el)


def essential(table,max_fields=5):
    fields=TABLES[table]['fields'];ordered=sorted(fields,key=lambda f:0 if f['pk'] else 1 if f['fk'] else 2)
    selected=ordered[:max_fields]
    return [f"{'PK ' if f['pk'] else 'FK ' if f['fk'] else ''}{f['name']} : {f['type'].replace('timestamptz(6)','fecha').replace('character varying','texto')}" for f in selected]


def case_classes(case,path):
    nodes=[{'name':case['boundary'],'attributes':['- '+x for x in UI_ATTRS[case['id']]],'operations':[]},
           {'name':case['control'],'attributes':['- contexto : ContextoTienda','- permiso : PoliticaAcceso'],'operations':[]}]
    keys={'boundary':nodes[0],'control':nodes[1]}
    for entity in case['entities']:
        attrs=essential(ALIASES[entity],4) if entity in ALIASES else LOCAL_ATTRIBUTES.get(entity,['id : UUID','version : entero','estado : texto'])
        node={'name':entity,'attributes':['- '+a for a in attrs],'operations':[]};nodes.append(node);keys[entity]=node
    for m in case['messages']:
        if m['to'] not in keys:continue
        method=re.sub(r'^\d+[.\s)]*','',m['text']);method=re.sub(r'\([^)]*\)','()',method)
        if not method.endswith(')'):
            tokens=re.findall(r'[A-Za-z0-9]+',unicodedata.normalize('NFKD',method).encode('ascii','ignore').decode())
            method=tokens[0].lower()+''.join(token[0].upper()+token[1:] for token in tokens[1:])+'()'
        # El contrato completo se conserva en el diagrama de secuencia.
        target=keys[m['to']]
        maximum=2 if m['to']=='control' else 1
        if '+ '+method not in target['operations'] and len(target['operations'])<maximum:target['operations'].append('+ '+method)
    for node in nodes:
        if not node['operations']:node['operations']=['+ validarEstado()']
    rel=[{'from':case['boundary'],'to':case['control'],'label':'solicitud validada'}]+[
        {'from':case['control'],'to':e,'label':'operación en el contexto autorizado'} for e in case['entities']]
    return graph_boxes(case['title'],nodes,rel,path)


def overview(book):
    book.toc()
    book.heading('1 Perfil del proyecto',newpage=True)
    book.heading('1.1 Introducción',2)
    book.para('Lúmina es una plataforma de comercio electrónico de moda femenina cuya función principal es el vestidor 3D con avatar personal. Esta segunda entrega define su evolución a un servicio SaaS para tiendas independientes, trabajo offline en web y Android, reportes interactivos y análisis mediante machine learning. Elaboramos primero esta especificación para desarrollar y verificar el sistema conforme a ella.')
    book.para('El informe aplica ICONIX: conecta requisitos y modelo de dominio con casos de uso, análisis de robustez, secuencias y clases de diseño. Incluye la base de datos, los contratos de sincronización, la evaluación de modelos y los criterios de aceptación. Los artefactos de diseño describen el sistema objetivo; los resultados de aceptación se registrarán después de ejecutar las pruebas durante el desarrollo.')
    book.heading('1.2 Objetivo general',2)
    book.para('Desarrollar una plataforma SaaS web y Android para tiendas de ropa femenina que combine vestidor 3D, operación comercial consistente, inventario offline con sincronización, dashboards personalizados y una demostración de analítica con modelos entrenados y evaluados.')
    book.heading('1.3 Objetivos específicos',2)
    book.numbered([
        'Separar cuentas, catálogo, inventario, pedidos, archivos y reportes de cada tienda, con tres planes mensuales y permisos por ubicación.',
        'Generar un avatar aproximado a partir de tres fotografías de cuerpo completo y altura declarada; permitir su revisión, corrección, eliminación y uso en el vestidor 3D web.',
        'Conservar catálogo, compra web/Android, mostrador y gestión de pedidos y devoluciones con movimientos transaccionales de existencias.',
        'Permitir a las compradoras descargar hasta diez prendas compatibles y a las administradoras crear, editar, archivar, ajustar y contar inventario sin conexión.',
        'Construir reportes con filtros avanzados y dashboards interactivos, también sobre conjuntos previamente descargados; interpretar texto o voz cuando exista conexión.',
        'Demostrar entrenamiento y evaluación de demanda y anomalías con datos sintéticos identificados, sin presentarlos como ventas ni exactitud comercial real.',
        'Mantener trazabilidad entre requisitos, casos, diagramas, datos, módulos y pruebas planificadas.'
    ])
    book.heading('1.4 Problema y alcance',2,newpage=True)
    book.para('Una tienda necesita integrar la consulta de prendas, las pruebas virtuales, las ventas por distintos canales y las existencias de sucursales y almacenes. Las interrupciones de internet no deben impedir el conteo ni la captura de cambios administrativos. Al crecer a varias tiendas independientes, compartir una sesión o una caché sin ámbito puede mezclar datos; los saldos offline tampoco pueden sustituir sin validación el inventario que ya recibió ventas.')
    book.para('La plataforma atiende compradoras, administradoras, vendedoras, analistas y una operadora de plataforma. El catálogo comercial es femenino; el personal no está restringido por género. La operadora prepara los espacios SaaS y administra planes, mientras el equipo de una tienda solo trabaja sobre su propio negocio y ubicaciones autorizadas.')
    book.para('La web utiliza React, TypeScript y Three.js; la API NestJS y Prisma con PostgreSQL; Android React Native y Expo; el procesamiento corporal Python, MediaPipe y Blender. La extensión ML utiliza Python y scikit-learn. Los componentes se organizan por módulos de negocio y contratos comunes, conservando la base existente.')
    book.heading('1.5 Límites de la experiencia virtual',2)
    book.para('El avatar representa una aproximación paramétrica, no una reconstrucción fotorealista ni una medición clínica. Las mallas de las prendas deben declarar plantilla, versión y licencia; una silueta referencial de vestido o falda no reproduce necesariamente el corte exacto del producto de marca. La interfaz no promete talla ni ajuste físico garantizados.')
    book.para('En web se muestran modelos GLB sobre avatar compatible. En Android la experiencia AR emplea imágenes transparentes de prendas ancladas a la pose detectada localmente; se identifica como superposición aproximada. Un PNG no se denomina modelo 3D. La generación del cuerpo desde fotografías requiere servidor; probar recursos ya descargados puede hacerse offline.')
    book.heading('1.6 Línea base y control documental',2)
    book.para('Versión de diseño 2.0, aprobaciones de alcance al 31 de octubre de 2026; entrega prevista para el 7 de noviembre de 2026. Repositorio oficial Santyxd353/examenSI2REPO. La rama desarrollo/saas-offline contiene la base de aislamiento SaaS y los módulos web; esta especificación define las ampliaciones antes de implementarlas.')
    book.para('El historial del incremento 6fcac5a, fechado el 20 de octubre, registra 100 pruebas de API y verificaciones web y móvil. Ese antecedente no acredita los nuevos casos offline, cobros, dashboards o ML. El capítulo de pruebas contiene criterios y escenarios planificados, sin resultados de ejecución inventados. La matriz de alcance se actualizará tras cada incremento.')


def plans_and_rules(book):
    book.heading('2 Requisitos y reglas de negocio',newpage=True)
    book.heading('2.1 Planes mensuales',2)
    plans=SPEC['plans']
    rows=[('Precio mensual','Bs 149','Bs 299','Bs 599'),('Ubicaciones activas','1','5','20'),('Personal activo','3','15','100'),('Productos registrados','200','2.000','10.000'),('Avatares generados por período','30','150','500'),('Reportes guardados','3','15','50'),('Interpretaciones IA por período','0','100','500'),('Machine learning','No incluido','Demanda y reposición','Demanda, anomalías y sucursales')]
    book.table(['Servicio','Esencial','Profesional','Negocio'],rows,[2.60,1.3,1.54,1.64])
    book.para('Los tres planes incluyen web y APK, vestidor, catálogo, inventario, ventas, pedidos, filtros manuales, dashboards offline sobre datos descargados y un paquete de hasta diez prendas por compradora. Personal incluye a la administradora y excluye a las compradoras. Una ubicación puede ser sucursal o almacén; una variante de talla/color no consume otro producto. Los productos archivados conservan historia y siguen contando como productos registrados.')
    book.para('Las cuotas de IA y generación de avatares se comparten por tienda y período, no por dispositivo. La visualización repetida de un avatar existente no es generación nueva. Al aceptar un trabajo se reserva cupo; se confirma con el resultado correcto o se libera si falla. Una regeneración solicitada y completada consume otra unidad. Reintentar el mismo trabajo o consulta no duplica el consumo.')
    book.para('El ciclo mensual es fecha a fecha según el período contratado; se almacenan instantes UTC y se presentan en America/La_Paz. La prueba inicial de catorce días no confirma un pago. El espacio Lúmina heredado es compatibilidad interna y no un cuarto plan comercial.')
    book.heading('2.2 Cobro y vigencia',2,newpage=True)
    book.para('La entrega académica integra Stripe en modo TEST y QR explícitamente simulado. El servidor fija precio, moneda BOB, tienda y plan: el navegador no decide el importe autorizado. Stripe se confirma con webhook de firma válida; el retorno a una pantalla de éxito no activa por sí mismo la suscripción. Los eventos se deduplican por proveedor e identificador externo y se comprueba su vigencia y orden.')
    book.para('Cada QR de demostración vincula tienda, plan, importe y período; su confirmación manual autorizada queda marcada QR_SIMULATED. No representa una transferencia bancaria. Para cobros comerciales reales se necesita una cuenta Stripe elegible; Bolivia no figura en la disponibilidad de cuentas Payments consultada al elaborar el diseño. El sistema conserva los precios en bolivianos sin fingir liquidaciones bancarias reales.')
    book.para('Un cambio inmediato de plan inicia un período completo con el importe y condiciones que la administradora acepta antes de confirmar; no se calcula un prorrateo implícito. Si el nuevo límite es menor que el uso, se rechaza el cambio hasta regularizarlo y no se borran registros. Cancelar la próxima renovación conserva la vigencia pagada hasta fin de período; luego se deriva EXPIRED y se habilita consulta sin nuevas escrituras comerciales.')
    book.table(['Estado persistido','Significado','Acceso'],[
        ['PROVISIONING / FAILED','Preparación en curso o fallida','Sin acceso comercial'],
        ['TRIAL + TRIALING','Tienda preparada en prueba','Prestaciones de prueba hasta su vencimiento'],
        ['ACTIVE + ACTIVE','Suscripción vigente','Según plan y permisos'],
        ['ACTIVE + EXPIRED derivado','Fin de prueba o período','Lectura y autenticación; renovación disponible'],
        ['SUSPENDED','Suspensión administrativa','Sin entregar espacio; mantenimiento de privacidad activo'],
    ],[2.05,2.9,2.13])
    book.heading('2.3 Actores y módulos',2,newpage=True)
    book.table(['Actor','Responsabilidad','Ámbito'],[
      ['Compradora','Colección, perfil, avatar, prueba virtual, compra y devolución','Su cuenta y tienda'],
      ['Administradora','Catálogo, equipo, ubicaciones, inventario, pedidos, reportes y plan','Su negocio'],
      ['Vendedora','Stock y mostrador','Ubicaciones asignadas'],
      ['Analista','Reportes y análisis autorizados','Ubicaciones asignadas'],
      ['Operadora de plataforma','Provisión, planes y soporte','Registro global; sin acceso comercial implícito'],
      ['Stripe TEST','Eventos autenticados de suscripción','Proveedor externo'],
    ],[1.55,3.75,1.78])
    book.para('La administración se divide en Inicio, Catálogo y vestidor, Inventario, Sucursales y almacenes, Ventas, Pedidos y devoluciones, Equipo, Reportes y análisis, Sincronización y Mi negocio. La navegación oculta acciones no permitidas y el servidor verifica cada operación aunque se invoque directamente la API. Android amplía su dashboard con los módulos de catálogo, inventario y sincronización que exige el trabajo offline.')
    groups=defaultdict(list)
    for r in SPEC['requirements']:groups[r['module']].append(r)
    book.heading('2.4 Requisitos funcionales',2,newpage=True)
    for idx,(module,rs) in enumerate(groups.items()):
        if idx:book.heading(module,3)
        else:book.heading(module,3)
        rows=[]
        for r in rs:
            cases=', '.join(c['id'] for c in SPEC['cases'] if r['id'] in c['requirements'])
            rows.append([r['id'],r['text'],cases])
        book.table(['ID','Requisito verificable','Casos'],rows,[.57,5.24,1.27])
    book.heading('2.5 Requisitos de calidad',2,newpage=True)
    qualities=[
      ['RNF01','Aislamiento','Dos tiendas con el mismo correo y UUID no comparten datos, archivos, eventos, cookies ni cachés.'],
      ['RNF02','Seguridad','TLS, contraseñas hash, sesiones revocables y secretos solo en servidor; revalidación de permisos al sincronizar.'],
      ['RNF03','Durabilidad offline','Cierre y reapertura conservan intenciones pendientes; respuesta perdida no duplica el movimiento confirmado.'],
      ['RNF04','Consistencia','Stock disponible no negativo; cambios concurrentes respetan reservas, compromisos, versión e idempotencia.'],
      ['RNF05','Privacidad','Consentimiento para fotografías y copia local del avatar; purga de capturas al procesar y antes de 24 horas; eliminación trazable.'],
      ['RNF06','Accesibilidad','Contraste legible, controles etiquetados, navegación por teclado en web y errores junto a los campos.'],
      ['RNF07','Rendimiento','Snapshot paginado, caché explícita y carga diferida de modelos; medir tiempos y FPS en equipo y Android de prueba.'],
      ['RNF08','Explicabilidad','Reportes muestran cobertura, fecha y filtros; ML muestra procedencia, versión, métricas y limitaciones.'],
      ['RNF09','Reproducibilidad','Semilla, hash de dataset, parámetros, particiones temporales y versión del artefacto permiten repetir la demostración.'],
      ['RNF10','Mantenibilidad','Migraciones de todas las bases y contratos versionados; trazabilidad requisito a caso, clase y aceptación.'],
    ]
    book.table(['ID','Calidad','Criterio'],qualities,[.62,1.28,5.18])
    book.heading('2.6 Reglas transversales',2,newpage=True)
    rules=[
      'RB01 El token y el contexto de la petición deben identificar la misma tienda. Un identificador desconocido o suspendido no cae en Lúmina por defecto.',
      'RB02 La identidad comercial usa una base por tienda; no se crean claves foráneas ficticias entre bases de negocio y el registro de control.',
      'RB03 Provisión completa, pago confirmado y permiso de negocio son condiciones diferentes. Ninguna reemplaza a las otras.',
      'RB04 Capacidades y cuotas se comprueban en servidor bajo bloqueo y transacción; se cuenta una persona una vez aunque tenga varios roles.',
      'RB05 Se archivan prendas con historial; no se borran pedidos ni movimientos por retirar una publicación.',
      'RB06 Un conteo con versión vencida no sustituye el saldo. Un delta offline se aplica sobre el saldo del servidor validado en ese momento.',
      'RB07 Una intención tiene UUID y hash inmutables. Corregirla genera otra intención enlazada; reutilizar un UUID con otro contenido produce conflicto.',
      'RB08 No se confirman compras, pagos, transferencias entre ubicaciones ni generación de avatar sin conexión al servidor.',
      'RB09 Las fotos corporales son privadas, con consentimiento y retención acotada; el video AR se procesa localmente y no se envía al servidor.',
      'RB10 La IA produce una definición validable; los importes y cantidades provienen de agregaciones autorizadas, no del texto generado.',
      'RB11 Los datos sintéticos, modelos y salidas de demostración se identifican y se aíslan de ventas, stock, cobros y estadísticas reales.',
      'RB12 Predecir demanda o señalar una anomalía no ajusta stock, compra mercadería ni acusa fraude automáticamente.',
      'RB13 Los cambios de permisos y suspensión se aplican en servidor al reconectar; no se promete retirar instantáneamente un archivo de un equipo sin red.',
      'RB14 Se preservan intención pendiente y evidencia ante errores; una respuesta técnica fallida no se registra como sincronización correcta.',
    ]
    for r in rules:book.para(r)


def domain_and_usecases(book):
    book.heading('3 Modelo de dominio y casos de uso',newpage=True)
    book.para('El modelo de dominio identifica conceptos y relaciones del problema. Las clases de diseño posteriores asignan operaciones a fronteras, controles y entidades. Una entidad local de la aplicación, como PaqueteOffline, no implica otra tabla en el servidor. RecursoVestidor abstrae modelos GLB y recursos transparentes AR sin equiparar sus formatos.')
    domain_groups=[
      ('Negocio y acceso',[
       {'name':'Tienda','attributes':['id : UUID','slug : texto','estado : EstadoTienda']},
       {'name':'Suscripcion','attributes':['plan : Plan','estado : EstadoSuscripcion','periodo : fechas']},
       {'name':'Plan','attributes':['codigo : texto','precioBOB : decimal','limites : capacidades']},
       {'name':'Usuario','attributes':['id : UUID','correo : texto','roles : conjunto']},
      ],[{'from':'Tienda','to':'Suscripcion','label':'tiene su período y plan'},{'from':'Plan','to':'Suscripcion','label':'define precio y capacidades'},{'from':'Tienda','to':'Usuario','label':'usuarios en su propia base'}]),
      ('Operación y vestidor',[
       {'name':'Producto','attributes':['id : UUID','nombre : texto','estado : texto']},
       {'name':'Variante','attributes':['sku : texto','talla : texto','color : texto']},
       {'name':'Inventario','attributes':['fisico : entero','reservado : entero','version : entero']},
       {'name':'Avatar','attributes':['propietaria : Usuario','version : entero','plantilla : referencia']},
       {'name':'RecursoVestidor','attributes':['formato : GLB / PNG','variante : Variante','plantilla : referencia']},
      ],[{'from':'Producto','to':'Variante','label':'una o varias combinaciones'},{'from':'Variante','to':'Inventario','label':'existencias por ubicación'},{'from':'Variante','to':'RecursoVestidor','label':'recursos publicados y versionados'},{'from':'Avatar','to':'RecursoVestidor','label':'compatibilidad de plantilla; no propiedad'}]),
      ('Sincronización y conocimiento',[
       {'name':'OperacionOffline','attributes':['id : UUID','versionBase : entero','intencion : JSON']},
       {'name':'ComprobanteOperacion','attributes':['hash : SHA256','resultado : JSON','estado : terminal']},
       {'name':'DefinicionReporte','attributes':['filtros : JSON','metricas : catálogo','diseno : JSON']},
       {'name':'ConjuntoEntrenamiento','attributes':['origen : SINTETICO','semilla : entero','hash : SHA256']},
       {'name':'ModeloAnalitico','attributes':['algoritmo : texto','version : texto','metricas : resultado']},
      ],[{'from':'OperacionOffline','to':'ComprobanteOperacion','label':'confirmación idempotente'},{'from':'ConjuntoEntrenamiento','to':'ModeloAnalitico','label':'datos de entrenamiento evaluados'}]),
    ]
    for idx,(title,nodes,rels) in enumerate(domain_groups,1):
        book.heading(f'3.1.{idx} {title}',3,newpage=True)
        book.figure(graph_boxes(title,nodes,rels,ART/f'dominio-{idx}'),f'Modelo de dominio  {title}',maxheight=7.6)
    book.heading('3.2 Vista de casos de uso',2,newpage=True)
    packs=[('Plataforma y acceso',range(1,6)),('Catálogo e inventario',range(6,12)),('Vestidor y experiencia personal',range(12,16)),('Comercio',range(16,20)),('Trabajo offline',range(20,24)),('Reportes y aprendizaje',range(24,28))]
    for i,(title,ids) in enumerate(packs):
        if i:book.heading(title,3,newpage=True)
        else:book.heading(title,3)
        cases=[c for c in SPEC['cases'] if int(c['id'][2:]) in ids]
        book.figure(use_cases(cases,ART/f'casos-{i+1}',title),f'Actores y objetivos  {title}',maxheight=7.7)
    book.heading('3.3 Especificaciones de casos',2,newpage=True)
    book.para('Los pasos siguientes son contratos de comportamiento del sistema objetivo. Los flujos alternativos se enlazan con el número de paso que los origina. La autenticación y la comprobación de plan y permisos se reutilizan como condiciones comunes, sin convertir todos los objetivos en un caso genérico de gestión.')
    for case in SPEC['cases']:
        book.heading(case['id']+' '+case['title'],2,anchor=case['id'])
        book.para(case['actor']+' · '+case['module'],bold_label='Actor y módulo:')
        book.para(', '.join(case['requirements']),bold_label='Requisitos:')
        book.para(' '.join(case['preconditions']),bold_label='Precondiciones:')
        book.para(' '.join(case['postconditions']),bold_label='Postcondiciones:')
        book.heading('Flujo principal',3);book.numbered(case['main'])
        book.heading('Flujos alternativos',3)
        for alt in case['alternatives']:
            book.para(' '.join(alt['steps']),bold_label=f"Paso {alt['at']}  {alt['label']}:")
        book.heading('Reglas del caso',3)
        for rule in case['rules']:book.para(rule)
        book.heading('Criterios de aceptación',3)
        for item in case['acceptance']:book.para(item,keep=True)
        p=book.para('Diagramas asociados: ')
        for j,(name,prefix) in enumerate([('robustez','R'),('secuencia','S'),('clases','C')]):
            if j:p.add_run(' | ')
            book.link(p,name,anchor=f'{prefix}_{case["id"]}')


def case_diagrams(book):
    book.heading('4 Análisis de robustez',newpage=True)
    book.para('Cada frontera representa una pantalla o entrada al sistema; el control coordina una intención y comprueba las reglas; la entidad conserva información de dominio o estado local. La persona interactúa con la frontera, no directamente con la persistencia. Los diagramas muestran responsabilidades; las secuencias precisan mensajes y transacciones.')
    book.para('Convenciones: actor de palitos; frontera con barra lateral; control con marca direccional; entidad con subrayado. Las consultas y escrituras se ejecutan en la tienda autorizada. Los rótulos de entidades locales no crean por sí solos tablas de servidor.')
    for index,case in enumerate(SPEC['cases']):
        book.heading(case['id']+' '+case['title'],3,newpage=index%2==1,anchor='RB_'+case['id'])
        book.figure(robustness(case,ART/f'robustez-{case["id"]}'),case['id']+'  Análisis de robustez',anchor='R_'+case['id'],maxheight=3.95,width=6.72)
    book.heading('5 Diseño de secuencias',newpage=True)
    book.para('Los mensajes concretan la colaboración prevista para cada caso. Las líneas discontinuas representan respuestas hacia un participante anterior. Antes de una modificación se validan identidad, prestación contratada, permiso y versión. El detalle de parámetros y restricciones permanece en el caso y en los contratos del capítulo 8.')
    for index,case in enumerate(SPEC['cases']):
        book.heading(case['id']+' '+case['title'],2,newpage=index>0)
        image=sequence(case,ART/f'secuencia-{case["id"]}')
        book.figure(image,case['id']+'  Secuencia principal',anchor='S_'+case['id'],maxheight=7.7 if index else 7.0)
        book.para('Alternativas y resultado esperado: véase '+case['id']+'. Los proveedores externos y repositorios se invocan mediante el control; un retorno de interfaz no sustituye la confirmación transaccional.',small=True)
    book.heading('6 Clases de diseño',newpage=True)
    book.para('Se muestran las clases principales de cada colaboración con atributos y operaciones. Los nombres son contratos del diseño objetivo, no una afirmación de que todos los métodos ya existan. Las pantallas y servicios encapsulan entradas y reglas; las entidades se relacionan con el diccionario o con los almacenes locales declarados. Los diagramas presentan métodos principales y las secuencias conservan el contrato completo de mensajes.')
    for index,case in enumerate(SPEC['cases']):
        book.heading(case['id']+' '+case['title'],2,newpage=index>0)
        image=case_classes(case,ART/f'clases-{case["id"]}')
        book.figure(image,case['id']+'  Clases de diseño',anchor='C_'+case['id'],maxheight=7.75 if index else 7.0)


def data_chapter(book):
    book.heading('7 Modelo y diccionario de datos',newpage=True)
    book.heading('7.1 Criterio de persistencia',2)
    book.para('El esquema existente contiene 49 tablas comerciales y cuatro tablas de control platform. El diseño agrega seis tablas y amplía columnas de versiones, analítica y facturación. El diccionario distingue columnas existentes de las propuestas con la marca NUEVO; la presencia de una tabla heredada no implica una prestación nueva disponible en la interfaz.')
    book.para('El registro de control se consulta con una conexión fija; cada negocio comercial usa su propia base y directorio. La configuración actual permite alojar el control en la base inicial, en el esquema platform; para producción se configura su conexión explícita. No existe una clave foránea garantizada entre una tabla comercial y platform.tenant de otra base. El contexto del servidor y las credenciales de conexión establecen el ámbito.')
    book.para('Las entidades locales PaqueteOffline, RecursoLocal, OperacionOffline, ConflictoSincronizacion y DisenoDashboard se guardan en el cliente o como JSON del agregado correspondiente. ComprobanteOperacion se materializa en operacion_sincronizacion; ConjuntoEntrenamiento corresponde a conjunto_entrenamiento. RecursoVestidor reúne conceptualmente modelo_prenda y recurso_catalogo según formato y uso.')
    book.heading('7.2 Modelos lógicos por grupo',2,newpage=True)
    for index,group in enumerate(SPEC['data']['diagram_groups']):
        book.heading(group['id']+' '+group['title'],3,newpage=index>0)
        nodes=[{'name':name,'attributes':essential(name)} for name in group['tables']]
        book.figure(graph_boxes(group['title'],nodes,group['relations'],ART/f'datos-{group["id"]}',model=True),group['id']+'  Relaciones y cardinalidades',maxheight=7.65)
        refs=group.get('references',[])
        if refs:book.para('Referencias a otros grupos: '+', '.join(refs)+'.',small=True)
    book.heading('7.3 Diccionario completo',2,newpage=True)
    book.para('PK identifica una clave primaria y FK una referencia local; Nulo indica si se admite ausencia. UUID es un identificador de 128 bits. JSONB conserva estructuras con esquema validado por aplicación. Las columnas derivadas y restricciones CHECK de PostgreSQL se mantienen según las migraciones; una FK polimórfica no se inventa para entidades de distinto tipo.')
    planned_names={t['name'] for t in SPEC['data']['new_tables']}
    for table in SPEC['data']['existing_tables']+SPEC['data']['new_tables']:
        name=table['name'];book.heading('Tabla '+name,3)
        book.para('PLANIFICADA' if name in planned_names else 'EXISTENTE con ampliaciones identificadas cuando corresponda',small=True,keep=True)
        fields=TABLES[name]['fields']
        rows=[]
        for f in fields:
            key=f.get('key') or ('PK/FK' if f['pk'] and f['fk'] else 'PK' if f['pk'] else 'FK' if f['fk'] else '—')
            ref=f['fk'] or ('Identificador / valor del registro' if f['pk'] else 'Valor validado del atributo')
            if f.get('planned') or name in planned_names:ref='NUEVO · '+ref
            rows.append([f['name'],f['type'],'Sí' if f['nullable'] else 'No',key,ref])
        book.table(['Campo','Tipo','Nulo','Clave','Referencia y condición'],rows,[2.1,1.17,.44,.52,2.35],font=9)
        if table.get('uniques'):book.para('Unicidad adicional: '+str(table['uniques']),small=True)
        notes=table.get('purpose') or table.get('notes','')
        book.para(notes,small=True)
    book.heading('7.4 Cambios e integridad',2,newpage=True)
    for change in SPEC['data']['changes']:
        book.para(change['reason'],bold_label=change['table']+':')
    for rule in SPEC['data']['integrity_rules'][1:]:book.para(rule)
    book.heading('7.5 Tablas de compatibilidad',2)
    book.para('Los siguientes grupos se conservan para integridad del esquema heredado. Este documento no activa avisos publicitarios, notificaciones push ni búsqueda por imagen como prestaciones de la entrega.')
    for entry in SPEC['data']['legacy_scope']:book.para(entry['reason'],bold_label=', '.join(entry['tables'])+':')


def contracts(book):
    book.heading('8 Arquitectura y contratos técnicos',newpage=True)
    book.heading('8.1 Componentes y despliegue',2)
    nodes=[
      {'name':'Web React','attributes':['Three.js · IndexedDB','caché de aplicación','motor local de reportes']},
      {'name':'API NestJS','attributes':['contexto de tienda','permisos · cuotas · transacciones','contratos versionados']},
      {'name':'Android Expo','attributes':['AR local · SQLite','archivos privados','cola de operaciones']},
      {'name':'Persistencia PostgreSQL','attributes':['control platform fijo','una base por negocio','archivos por tienda']},
      {'name':'Procesos Python','attributes':['MediaPipe · Blender','RandomForest · IsolationForest','artefactos y métricas']},
    ]
    relations=[{'from':'Web React','to':'API NestJS','label':'HTTPS y token por tienda'},{'from':'Android Expo','to':'API NestJS','label':'HTTPS; sincronización al reconectar'},{'from':'API NestJS','to':'Persistencia PostgreSQL','label':'repositorios según contexto'},{'from':'API NestJS','to':'Procesos Python','label':'trabajos y resultados identificados'}]
    book.figure(graph_boxes('Arquitectura',nodes,relations,ART/'arquitectura'),'Arquitectura lógica y responsabilidades',maxheight=7.4)
    book.heading('8.2 Sincronización offline',2,newpage=True)
    book.para('IndexedDB en web y SQLite más archivos privados en Android conservan snapshots y una cola por tienda, persona e instalación. Un service worker permite reabrir la aplicación web descargada. La identidad local habilita trabajo sobre la copia descargada; no crea una sesión nueva de servidor ni sustituye su autorización. Un cierre de sesión bloquea la copia anterior y exige reautenticar a su propietaria antes de sincronizar.')
    book.table(['Campo de intención','Contrato'],[
      ['operationId','UUID estable; no cambia durante reintentos'],['deviceId','UUID de instalación, independiente de token push'],
      ['kind','CATALOG_CREATE / UPDATE / ARCHIVE, STOCK_DELTA o STOCK_COUNT'],['entityId','UUID del agregado; altas generan identificador antes de enviar'],
      ['baseVersion','Versión descargada; obligatoria en edición y conteo'],['payload','Solo campos del esquema de la operación'],['createdAt','Fecha local informativa, sin autoridad sobre el saldo'],
    ],[1.55,5.53])
    book.para('GET /api/offline/snapshot devuelve ámbito, versiones, fecha y cobertura paginada. POST /api/offline/sync recibe lotes acotados, con dependencias y resultado por intención. GET /api/offline/outfits devuelve un manifiesto de hasta diez prendas compatibles. El actor proviene del token y la tienda del contexto; el cuerpo no puede imponerlos.')
    book.para('La cola persiste PENDING, SENDING, APPLIED, CONFLICT y BLOCKED. Al reabrir tras un envío interrumpido, SENDING vuelve a pendiente para recuperar el comprobante. El servidor registra APPLIED, CONFLICT o REJECTED como resultados terminales; un fallo técnico hace rollback. Una falta transitoria de sesión, permiso o vigencia conserva el bloqueo hasta revalidación y no se congela como un rechazo definitivo.')
    book.para('La mutación, el movimiento, la auditoría y el comprobante se confirman en la misma transacción de negocio. Una respuesta perdida se resuelve recuperando ese resultado. Si UUID, actor, instalación, tipo o hash no coinciden, se rechaza la reutilización. Los conflictos conservan la intención original y los valores actuales; corregir crea otra intención enlazada. Las altas del catálogo se sincronizan antes de sus movimientos dependientes.')
    book.heading('8.3 Descargas y seguridad local',2,newpage=True)
    book.para('El manifiesto incluye formato, variante, plantilla, versión, tamaño, SHA256 y URL autorizada. Se eligen hasta diez prendas según gustos, talla e interacciones disponibles; con poco historial se informa que la selección usa reglas. La clienta autoriza expresamente la copia de su avatar aprobado. La descarga solo se marca completa cuando se verificaron sus archivos; un paquete parcial muestra qué recursos son utilizables.')
    book.para('La caché distingue tienda, propietaria y versión. Los archivos privados no entran en una caché pública compartida por URL. Al cambiar de negocio no se muestra el paquete ni la cola anterior. Al cerrar sesión con operaciones pendientes se informa y se permite conservarlas bloqueadas o descartarlas con una decisión explícita. Se usan dispositivos confiables, almacenamiento privado móvil y protección de datos sensibles; la revocación remota se aplica al reconectar.')
    book.para('Los deltas se verifican contra reservas y compromisos del servidor. Un conteo vencido exige conciliación, no sobrescritura. Las fotos nuevas quedan guardadas localmente hasta poder enviarlas y validarlas; archivar conserva el historial. Comprar, cobrar, transferir entre ubicaciones y generar un cuerpo desde fotos siguen requiriendo conexión.')
    book.heading('8.4 Definición de reportes y dashboard',2,newpage=True)
    book.para('Un reporte guardado conserva versión de esquema, nombre, fuente, filtros, métricas, dimensiones, agrupación temporal, orden y diseño de paneles. Los filtros comprenden período y comparación, sucursal/almacén, canal, vendedora, categoría, marca, producto, talla, color, estado de pedido, método/estado del pago, devoluciones, precio y stock. El catálogo de campos, operadores y agregaciones es cerrado y se valida en servidor.')
    book.para('Se separa el grano de cálculo: ventas por pedido para contar tickets y total cobrado; detalle para unidades y desglose por prenda; pagos y devoluciones se preagregan para evitar multiplicar importes al unir varias filas. Se declara si una métrica es venta bruta confirmada, devolución o venta neta. Los precios históricos provienen del detalle de pedido, no del precio actual del catálogo.')
    book.table(['Interacción','Comportamiento verificable'],[
      ['Filtros y selección cruzada','Elegir barra, punto o fila aplica filtros visibles y permite restablecerlos'],
      ['Indicadores y detalle','Tickets, unidades, ingresos, devoluciones y ticket promedio respetan el mismo ámbito'],
      ['Diseño guardado','Tarjetas, serie temporal, barras y tabla se reabren con sus definiciones'],
      ['Offline','El mismo motor lógico filtra y agrega únicamente filas descargadas; muestra cobertura y fecha'],
      ['Exportación','CSV y vista PDF conservan filtros, fuente, fecha y marca de demostración cuando corresponda'],
      ['Sin filas','Muestra cero o sin datos según métrica; no crea ventas ni simula una descarga completa'],
    ],[1.7,5.38])
    book.heading('8.5 Asistencia por texto y voz',2)
    book.para('Con conexión, la voz se transcribe y el texto se interpreta como una definición de reporte permitida. La persona revisa filtros y alcance antes de ejecutar. Se valida plan, cuota y permiso; la IA no recibe credenciales ni genera SQL ejecutable. El servidor calcula los resultados y los presenta en el dashboard. Con internet ausente o proveedor no configurado quedan los filtros manuales, sin respuestas inventadas.')
    book.para('La integración usa un adaptador configurable de proveedor para interpretación y transcripción. El proveedor y las credenciales se fijarán durante la configuración del servicio, sin alterar el contrato de salida. Las solicitudes se miden por interpretación procesada una sola vez; ejecutar filtros manuales o volver a explorar el resultado no gasta otra unidad. No se conserva audio innecesario.')


def ml_chapter(book):
    book.heading('9 Protocolo de machine learning',newpage=True)
    book.heading('9.1 Demostración sintética identificada',2)
    book.para('La demostración usará un conjunto reproducible generado con semilla 42: 365 fechas del 1 de octubre de 2025 al 30 de septiembre de 2026, doce prendas femeninas, 36 variantes, cinco ubicaciones y tres canales. El panel diario completo tiene 197.100 combinaciones fecha-variante-ubicación-canal, incluidas cantidades cero. La generación incorpora calendario, preferencias y variación aleatoria, con anomalías etiquetadas exclusivamente para evaluar la demostración.')
    book.para('Los datos se guardan como conjunto de entrenamiento con namespace sintético, metadatos y hash. No se insertan pedidos, pagos, clientes ni saldos ficticios en las tablas comerciales reales. El laboratorio, sus gráficos y exportaciones llevan la marca DEMOSTRACIÓN CON DATOS SINTÉTICOS. Sus métricas muestran aprendizaje sobre un problema generado, no precisión comprobada para una tienda real.')
    book.heading('9.2 Modelo de demanda',2)
    book.para('Se propone RandomForestRegressor con un pipeline de preparación: identificadores categóricos de variante, ubicación y canal; calendario; rezagos de 1, 7, 14 y 28 días; medias móviles anteriores de 7 y 28 días. Ninguna variable utiliza ventas de la fecha objetivo o de fechas posteriores. Las primeras 28 fechas sirven como calentamiento de características y no se evalúan como pronósticos sin historia.')
    book.para('El modelo produce demanda a 7 y 30 días. Para cada origen de pronóstico solo se observa la historia disponible hasta ese corte; los pasos futuros se calculan recursivamente y alimentan sus propias predicciones, sin tomar las ventas reales futuras como rezagos. Se guarda la versión del pipeline, parámetros, semilla, datos y artefacto entrenado.')
    book.heading('9.3 Evaluación temporal',2,newpage=True)
    book.table(['Partición por fecha','Período','Uso'],[
      ['Entrenamiento · 255 fechas','01/10/2025–12/06/2026','Ajustar transformaciones y modelo con datos previos'],
      ['Validación · 55 fechas','13/06/2026–06/08/2026','Elegir parámetros y comparar configuraciones'],
      ['Prueba · 55 fechas','07/08/2026–30/09/2026','Evaluación final sin elegir parámetros sobre sus resultados'],
    ],[2.0,2.1,2.98])
    book.para('El reparto es por fechas para todas las series, no por filas aleatorias. Tras seleccionar la configuración se puede reajustar con entrenamiento y validación, manteniendo intacta la prueba final. Los cortes walk-forward respetan el horizonte: solo se evalúan orígenes con futuro completo dentro de su bloque; la actualización de historia usa observaciones conocidas hasta cada nuevo origen.')
    book.para('La referencia es ingenua estacional: repetir la cantidad del mismo día de la semana previa, sin mirar el futuro. Se registran MAE, RMSE y WAPE para modelo y referencia, por horizonte y ámbito. WAPE no se calcula cuando la suma real es cero; se informa no aplicable. No se exige ni se inventa una mejora del bosque sobre la referencia: se muestra la comparación observada y se explica si no aporta ventaja.')
    book.heading('9.4 Anomalías y reposición',2)
    book.para('Negocio incorpora IsolationForest sobre características normalizadas y obtenidas con historia anterior. Las etiquetas de anomalías inyectadas se reservan para medir precisión, exhaustividad y F1; no se incluyen como variable de entrada. La puntuación identifica observaciones inusuales para revisión, sin afirmar fraude ni causa probada.')
    book.para('La reposición combina pronóstico, disponible, stock de seguridad y plazo de abastecimiento configurado. Un traslado sugerido solo compara ubicaciones de la misma tienda y exige existencias suficientes en origen. Son propuestas para decisión de la administradora; no modifican inventario ni crean compras. Profesional recibe demanda/reposición; Negocio agrega anomalías y comparación por ubicaciones.')
    book.heading('9.5 Evidencia que debe producir el sistema',2)
    book.para('Cada entrenamiento debe entregar dataset/hash, particiones, variables, parámetros, versión, duración y métricas. Cada consulta muestra horizonte, fecha de corte, procedencia, referencia de comparación y limitaciones. Si no existen artefacto válido, historial o métrica aplicable, se informa la limitación. Las proyecciones estadísticas heredadas se identifican como regla o referencia y no se renombran como un modelo aprendido.')


def testing_trace(book):
    book.heading('10 Pruebas y trazabilidad',newpage=True)
    book.heading('10.1 Estrategia de verificación',2)
    book.para('Las pruebas se ejecutarán al avanzar el sistema definido. Cada escenario parte de datos y roles explícitos; registra fecha, versión, resultado obtenido y evidencia. En esta especificación se documentan criterios de aceptación y pruebas previstas, sin declarar como aprobados los flujos nuevos. La verificación documental comprueba correspondencia de requisitos, casos, diagramas y datos.')
    common=[
      ['PA01','CU01–CU03','Dos tiendas con mismos UUID y correo','Tokens, archivos, SQL, sesiones y eventos no cruzan de negocio'],
      ['PA02','CU02','Webhook repetido o fuera de orden','Un solo efecto; no revive un período anterior'],
      ['PA03','CU02','QR de demostración','Importe/tienda/plan correctos y marca explícita de simulación'],
      ['PA04','CU02–CU06','Altas concurrentes cerca del límite','No supera puestos, productos ni ubicaciones'],
      ['PA05','CU09–CU11','Dos operaciones compiten con reservas','Saldo válido, bloqueo y resultado sin duplicar movimiento'],
      ['PA06','CU12–CU15','Fotos inválidas o recurso incompatible','Rechazo explicable, purga y ausencia de promesa de ajuste'],
      ['PA07','CU16–CU17','Doble confirmación o respuesta perdida','Un pedido, un pago y un movimiento por intención'],
      ['PA08','CU18–CU19','Devolución o cancelación repetida','Estado, reintegro y reembolso consistentes e idempotentes'],
      ['PA09','CU20–CU23','Cerrar y reabrir sin red','Paquetes/copia y pendientes durables dentro de su ámbito'],
      ['PA10','CU21–CU23','Conteo vencido, permiso revocado o plan vencido','Conflicto/bloqueo visible; no sobrescribe stock ni pierde intención'],
      ['PA11','CU24–CU25','Mismos filtros online y sobre copia completa','Mismos agregados y ámbito; copia parcial declara cobertura'],
      ['PA12','CU26–CU27','Entrenamiento y evaluación sintéticos','Semilla/hash reproducibles, separación temporal y etiqueta demo'],
    ]
    book.table(['ID','Casos','Escenario','Resultado esperado'],common,[.6,1.05,2.55,2.88],font=9)
    book.heading('10.2 Pruebas específicas del incremento',2,newpage=True)
    for i,test in enumerate(SPEC['acceptance_tests'],13):
        scenario=test.get('scenario',test.get('name','')).rstrip('. :')
        book.para(str(test.get('expected','')),bold_label=f'PA{i:02}  {test.get("case",test.get("cases",""))}  {scenario}:')
    book.heading('10.3 Matriz de trazabilidad',2,newpage=True)
    book.para('Cada caso conserva un trío de diagramas propio y criterios de aceptación en su especificación. Las columnas R, S y C enlazan robustez, secuencia y clases. Los criterios anteriores completan escenarios transversales y de concurrencia.')
    rows=[]
    for c in SPEC['cases']:rows.append([c['id'],', '.join(c['requirements']),c['module'],f'R / S / C de {c["id"]}', 'Planificada'])
    book.table(['Caso','Requisitos','Módulo','Diseño','Aceptación'],rows,[.6,1.42,2.18,1.63,1.25],font=9)
    book.heading('10.4 Control de cobertura',2)
    book.para('La especificación comprende 54 requisitos funcionales, 27 casos y sus diagramas de robustez, secuencia y clases. El diccionario cubre las 53 tablas existentes entre negocio y control, más seis propuestas. Los cambios de columnas y cuotas se distinguen de estructuras ya migradas. Toda función agregada durante el desarrollo debe actualizar sus requisitos, diseño y aceptación antes del cierre de la entrega.')


def closing(book):
    book.heading('11 Organización del desarrollo',newpage=True)
    book.para('El documento se prepara primero y se utiliza como contrato de implementación. El orden respeta dependencias: no se sincroniza sin aislamiento, no se evalúa ML sin conjunto versionado y no se declara dashboard offline sin filas descargadas y motor local de filtros.')
    book.table(['Fecha','Entregable','Condición para avanzar'],[
      ['31/10','Documento ICONIX y línea base de alcance','Coherencia de casos, planes, datos y diagramas'],
      ['01/11','Planes, prestaciones, cuotas y suscripciones TEST','Dos tiendas aisladas y cobro idempotente de demostración'],
      ['02/11','Protocolo y operación offline web','Cierre/reapertura, respuesta perdida y conflicto de versión'],
      ['03/11','Offline Android y paquete de prendas','Cola durable y recursos verificados utilizables sin red'],
      ['04/11','Constructor y dashboards web/Android','Filtros, detalle, exportación y cobertura offline verificables'],
      ['05/11','Texto/voz y laboratorio de ML sintético','Definición validada y modelos con métricas temporales'],
      ['06/11','Integración y APK firmado','Pruebas de dos tiendas y actualización compatible'],
      ['07/11','Verificación del despliegue y cierre documental','Evidencia real, trazabilidad y correcciones finales'],
    ],[.72,3.00,3.36])
    book.para('Cada bloque conserva un resultado ejecutable y pruebas adecuadas. Las prioridades no justifican completar diagramas con resultados imaginados. Si una capacidad no supera aceptación, se registra como pendiente y no se anuncia como implementada; cualquier cambio de alcance requiere actualizar la línea base documental.')
    book.heading('12 Conclusiones y recomendaciones',newpage=True)
    book.heading('12.1 Conclusiones',2)
    book.para('El diseño integra el vestidor como elemento central de una plataforma SaaS de moda femenina. La separación por tienda evita mezclar el negocio de distintas administradoras; el protocolo offline conserva intenciones y versiones en lugar de sustituir saldos. La definición común de reportes permite explorar datos descargados con el mismo significado que en línea.')
    book.para('La demostración ML incorpora entrenamiento, artefactos y evaluación temporal identificados. Su utilidad académica consiste en comprobar el proceso completo y sus límites, sin atribuir a datos sintéticos resultados comerciales reales. ICONIX vincula estas decisiones con casos concretos y colaboraciones de diseño que guiarán el código y sus pruebas.')
    book.heading('12.2 Recomendaciones',2)
    recommendations=[
      'Desarrollar según los casos y contratos versionados; revisar el documento junto con cada cambio de alcance.',
      'Respaldar el registro de control, todas las bases de negocio y sus archivos antes de desplegar; conservar una versión recuperable.',
      'Mantener secretos, firmas y cuentas de proveedor fuera del repositorio y de los clientes web/Android.',
      'Probar pérdida de conexión, cierre de aplicación, cambios de permisos y operaciones simultáneas antes de afirmar soporte offline.',
      'Validar visualmente avatar y prendas con fotografías consentidas y dispositivos reales; informar siempre que la experiencia es aproximada.',
      'Conservar el laboratorio sintético separado. Para una futura evaluación comercial se necesita historial autorizado y suficiente de la tienda.',
      'Medir uso de procesamiento y almacenamiento para revisar los cupos y precios antes de una comercialización real.',
    ]
    book.numbered(recommendations)
    book.heading('13 Bibliografía',newpage=True)
    refs=[
      ('Rosenberg, D. y Stephens, M. (2007). Use Case Driven Object Modeling with UML: Theory and Practice. Apress.','https://link.springer.com/book/10.1007/978-1-4302-0369-8'),
      ('Microsoft. SaaS and multitenant solution architecture.','https://learn.microsoft.com/en-us/azure/architecture/guide/saas-multitenant-solution-architecture/'),
      ('Microsoft. Tenancy models for a multitenant solution.','https://learn.microsoft.com/en-us/azure/architecture/guide/multitenant/considerations/tenancy-models'),
      ('MDN. IndexedDB API.','https://developer.mozilla.org/en-US/docs/Web/API/IndexedDB_API'),
      ('MDN. Service Worker API.','https://developer.mozilla.org/en-US/docs/Web/API/Service_Worker_API'),
      ('Expo. SQLite.','https://docs.expo.dev/versions/latest/sdk/sqlite/'),
      ('Stripe. Build a subscriptions integration.','https://docs.stripe.com/billing/subscriptions/build-subscriptions'),
      ('Stripe. Webhook endpoints y verificación de firmas.','https://docs.stripe.com/webhooks'),
      ('Stripe. Global availability.','https://stripe.com/global'),
      ('scikit-learn. RandomForestRegressor.','https://scikit-learn.org/stable/modules/generated/sklearn.ensemble.RandomForestRegressor.html'),
      ('scikit-learn. IsolationForest.','https://scikit-learn.org/stable/modules/generated/sklearn.ensemble.IsolationForest.html'),
      ('scikit-learn. TimeSeriesSplit.','https://scikit-learn.org/stable/modules/generated/sklearn.model_selection.TimeSeriesSplit.html'),
      ('Grupo 18. Repositorio de Lúmina y esquema de datos del proyecto.','https://github.com/Santyxd353/examenSI2REPO'),
    ]
    for text,url in refs:
        p=book.para(text);book.link(p,'  '+url,url=url)
    book.para('Fuentes técnicas consultadas el 31 de octubre de 2026. Las decisiones de cupos, contratos y alcance corresponden al diseño del proyecto; no son garantías de capacidad de los proveedores.',small=True)
    book.heading('14 Accesos al proyecto',newpage=True)
    book.para('Los códigos llevan al repositorio oficial y a la instancia publicada. El repositorio distingue el diseño y la rama de trabajo; la página desplegada puede reflejar la versión anterior hasta completar y verificar la nueva entrega.')
    for slug,title,url in [('repositorio','Repositorio oficial',SPEC['repository']),('despliegue','Página desplegada',SPEC['deployment'])]:
        book.heading(title,2)
        path=qr_code(url,ART/f'qr-{slug}.png')
        p=book.doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(path),width=Inches(1.85))
        p=book.para();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;book.link(p,url,url=url)


def main():
    ART.mkdir(parents=True,exist_ok=True)
    book=Book()
    overview(book);plans_and_rules(book);domain_and_usecases(book);case_diagrams(book);data_chapter(book);contracts(book);ml_chapter(book);testing_trace(book);closing(book)
    # Quitar relaciones de imágenes antiguas que ya no aparecen en el cuerpo.
    used=set(book.doc._element.xpath('//@r:embed'))
    for rel_id,rel in list(book.doc.part.rels.items()):
        if rel.reltype.endswith('/image') and rel_id not in used:book.doc.part.drop_rel(rel_id)
    book.doc.core_properties.title=SPEC['title'];book.doc.core_properties.author='Grupo 18'
    book.doc.core_properties.subject='Especificación ICONIX para desarrollo SaaS offline y analítica'
    book.doc.core_properties.keywords='ICONIX, Lúmina, Grupo 18, SaaS, offline, vestidor 3D'
    book.doc.save(OUTPUT)
    editable=BASE/'diagramas';editable.mkdir(parents=True,exist_ok=True)
    for path in book.figure_paths:
        svg=path.with_suffix('.svg')
        if svg.exists():copy2(svg,editable/svg.name)
    assert sha256(SOURCE.read_bytes()).hexdigest()==book.original_sha,'Se alteró la referencia'
    manifest={'docx':str(OUTPUT),'figures':book.counter,'requirements':len(SPEC['requirements']),'cases':len(SPEC['cases']),'tables':len(TABLES),'reference_sha256':book.original_sha}
    (ROOT/'.local/iconix/build-manifest.json').write_text(json.dumps(manifest,indent=2),encoding='utf8')
    print(json.dumps(manifest,ensure_ascii=False))


if __name__=='__main__':main()
